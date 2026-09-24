"""
Excel Data -> AI Prompt (step 1 of the pipeline).

Reads an uploaded .xlsx sheet of leads and turns each row into a structured
Lead object. Column names are matched case-insensitively and a handful of
common synonyms are supported so the sheet does not have to be pixel-perfect.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from .schemas import LeadCreate, LeadUploadSummary, Lead

# Maps accepted header spellings -> canonical field name.
HEADER_ALIASES: dict[str, str] = {
    "name": "name",
    "lead name": "name",
    "customer name": "name",
    "mobile": "mobile",
    "phone": "mobile",
    "phone number": "mobile",
    "mobile number": "mobile",
    "contact": "mobile",
    "contact number": "mobile",
    "service": "service",
    "service required": "service",
    "service interested in": "service",
    "budget": "budget",
    "budget (inr)": "budget",
    "city": "city",
    "location": "city",
}

REQUIRED_FIELDS = {"name", "mobile"}


def _normalize_header(raw: Any) -> str | None:
    if raw is None:
        return None
    key = str(raw).strip().lower()
    return HEADER_ALIASES.get(key)


def parse_leads_excel(file_bytes: bytes) -> LeadUploadSummary:
    """Parse an uploaded Excel file into validated Lead records.

    Expected columns (any order, case-insensitive): Name, Service, Budget, City,
    plus a mobile/phone column so the AI agent knows who to call.
    """

    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return LeadUploadSummary(total_rows=0, imported=0, skipped=0, errors=["Sheet is empty."])

    header_row = rows[0]
    field_by_column = {idx: _normalize_header(value) for idx, value in enumerate(header_row)}

    if not REQUIRED_FIELDS.issubset(set(field_by_column.values())):
        missing = REQUIRED_FIELDS - set(field_by_column.values())
        return LeadUploadSummary(
            total_rows=len(rows) - 1,
            imported=0,
            skipped=len(rows) - 1,
            errors=[f"Missing required column(s): {', '.join(sorted(missing))}"],
        )

    imported: list[Lead] = []
    errors: list[str] = []
    data_rows = rows[1:]

    for row_number, row in enumerate(data_rows, start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue

        record: dict[str, Any] = {}
        for idx, cell in enumerate(row):
            field = field_by_column.get(idx)
            if field and cell not in (None, ""):
                record[field] = str(cell).strip()

        try:
            lead = Lead(**LeadCreate(**record).model_dump())
            imported.append(lead)
        except Exception as exc:  # noqa: BLE001 - surfaced back to the caller
            errors.append(f"Row {row_number}: {exc}")

    return LeadUploadSummary(
        total_rows=len(data_rows),
        imported=len(imported),
        skipped=len(data_rows) - len(imported),
        errors=errors,
        leads=imported,
    )
